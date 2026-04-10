import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()

# Sheet 1: General Ledger
ws1 = wb.active
ws1.title = "1. General Ledger"
headers = ["Date", "Description", "Ref/Invoice", "Account", "Debit", "Credit", "Notes"]
ws1.append(headers)
for col in range(1, 8):
    ws1.cell(row=1, column=col).font = Font(bold=True)

ws1.append(["2026-01-01", "Clear Phantom Inventory", "ADJ-01", "Retained earnings", 451514.62, 0, "Wipes fake asset"])
ws1.append(["2026-01-01", "Clear Phantom Inventory", "ADJ-01", "Negative inventory clearing", 0, 451514.62, "Wipes fake asset"])
ws1.append(["2026-01-01", "Fix StadiumX Investment", "ADJ-02", "Equity investment", 9820.00, 0, "Moves from expense to asset"])
ws1.append(["2026-01-01", "Fix StadiumX Investment", "ADJ-02", "Retained earnings", 0, 9820.00, "Moves from expense to asset"])
ws1.append(["2026-01-01", "Clear Negative Tax Liab", "ADJ-03", "Profit tax 20%", 21623.78, 0, "Clears negative balances"])
ws1.append(["2026-01-01", "Clear Negative Tax Liab", "ADJ-03", "PPT 1%", 4928.79, 0, "Clears negative balances"])
ws1.append(["2026-01-01", "Clear Negative Tax Liab", "ADJ-03", "Retained earnings", 0, 26552.57, "Clears negative balances"])

# Sheet 2: The Missing $70k
ws2 = wb.create_sheet(title="2. Missing 70k Drop")
ws2.append(["Date", "Description / What was bought?", "Vendor / Paid To", "Amount (USD)", "Category (Inventory, Salary, etc.)"])
for col in range(1, 6):
    ws2.cell(row=1, column=col).font = Font(bold=True)

wb.save("/home/KOOMPI/.openclaw/nimmit/KOOMPI_2026_Clean_Ledger.xlsx")
