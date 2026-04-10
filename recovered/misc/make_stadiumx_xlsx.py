import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()

# Sheet 1: General Ledger
ws1 = wb.active
ws1.title = "1. General Ledger"
headers = ["Date", "Description", "Ref/Invoice", "Account", "Debit", "Credit", "Notes"]
ws1.append(headers)
for col in range(1, 8):
    ws1.cell(row=1, column=col).font = Font(bold=True)

# Adjusting Entry for StadiumX (Clear Employee Clearing Account)
ws1.append(["2026-01-01", "Clear Employee Clearing Account", "ADJ-01", "Employee clearing account", 5852.68, 0, "Clear liability (assign to actual expense or owner drawings)"])
ws1.append(["2026-01-01", "Clear Employee Clearing Account", "ADJ-01", "Retained earnings (or Expense Account)", 0, 5852.68, "Offset for cleared liability"])

# Sheet 2: New Transactions Drop
ws2 = wb.create_sheet(title="2. New Transactions Drop")
ws2.append(["Date", "Description / What was bought or sold?", "Vendor / Customer", "Amount (USD)", "Category (Income/Expense/Asset/Liability)"])
for col in range(1, 6):
    ws2.cell(row=1, column=col).font = Font(bold=True)

wb.save("/home/KOOMPI/.openclaw/nimmit/StadiumX_2026_Clean_Ledger.xlsx")
